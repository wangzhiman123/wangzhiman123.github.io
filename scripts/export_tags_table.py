#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
export_tags_table.py —— 把全部文章的分类、标签导出成一张表格

用途
----
把 _posts 里每篇文章开头的 categories / tags 抽出来，汇总成 tools\文章分类标签表.csv，
方便在 Excel / WPS 里统一查看、筛选、批量修改。

用法
----
    python scripts/export_tags_table.py

输出
----
    tools/文章分类标签表.csv       ← 主表格（UTF-8 带 BOM，Excel 双击不乱码）
    tools/文章分类标签表.xlsx      ← 只有在装了 openpyxl 时才会额外生成（可选）

表格列的含义
------------
    序号        仅用于查看，随便改
    一级分类    必须是 工作 / 生活 / 学习 之一
    二级分类    该大类下的细分，可自由填写（新分类会自动生成分类页）
    标签        多个标签用顿号「、」隔开（中文逗号、英文逗号也认）
    文章标题    只读参考，脚本不会用它，改了也不影响
    文件名      关键！脚本靠这一列找文章，千万不要改
    发布日期    只读参考

注意
----
这个脚本只「读」文章，不会修改任何文章内容。导出的表格会覆盖上一次的旧表格。
"""

import csv
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
POSTS_DIR = ROOT / "_posts"
OUT_DIR = ROOT / "tools"

CSV_NAME = "文章分类标签表.csv"

# 一级分类的推荐顺序（导出排序用，不是限制）
LEVEL1_ORDER = ["工作", "生活", "学习"]

RE_TITLE = re.compile(r"^title\s*:\s*(.+?)\s*$", re.M)
RE_DATE = re.compile(r"^date\s*:\s*(.+?)\s*$", re.M)


def read_list_field(text: str, key: str):
    """读 categories / tags 这类「方括号列表」字段，返回字符串列表（读不到返回 None）。"""
    m = re.search(r"^" + key + r"\s*:\s*\[(.*?)\]", text, re.M)
    if not m:
        return None
    inner = m.group(1).strip()
    if not inner:
        return []
    return [x.strip().strip('"').strip("'") for x in inner.split(",") if x.strip()]


def read_scalar_field(text: str, pattern, default: str = "") -> str:
    m = pattern.search(text[:4000])
    if not m:
        return default
    return m.group(1).strip().strip('"').strip("'")


def collect():
    """扫描全部文章，返回 [{序号, 一级分类, 二级分类, 标签, 文章标题, 文件名, 发布日期}, ...]"""
    rows = []
    for path in sorted(POSTS_DIR.glob("*.md")):
        text = path.read_text(encoding="utf-8", errors="replace")
        cats = read_list_field(text, "categories") or []
        tags = read_list_field(text, "tags") or []
        rows.append({
            "一级分类": cats[0] if len(cats) > 0 else "",
            "二级分类": cats[1] if len(cats) > 1 else "",
            "标签": "、".join(tags),
            "文章标题": read_scalar_field(text, RE_TITLE, path.stem),
            "文件名": path.name,
            "发布日期": read_scalar_field(text, RE_DATE, "")[:10],
        })

    def sort_key(r):
        c1 = r["一级分类"]
        idx = LEVEL1_ORDER.index(c1) if c1 in LEVEL1_ORDER else len(LEVEL1_ORDER)
        return (idx, c1, r["二级分类"], r["文件名"])

    rows.sort(key=sort_key)
    for i, r in enumerate(rows, 1):
        r["序号"] = i
    return rows


def write_csv(rows, path: Path):
    header = ["序号", "一级分类", "二级分类", "标签", "文章标题", "文件名", "发布日期"]
    # utf-8-sig：带 BOM，Excel 双击打开中文不乱码
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header, quoting=csv.QUOTE_ALL)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in header})


def write_xlsx(rows, path: Path):
    """装了 openpyxl 才走这里；没装就静默跳过，不影响主流程。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "分类标签"
    header = ["序号", "一级分类", "二级分类", "标签", "文章标题", "文件名", "发布日期"]
    ws.append(header)

    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([r.get(k, "") for k in header])

    widths = {"A": 6, "B": 12, "C": 14, "D": 40, "E": 50, "F": 30, "G": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column_letter == "D"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return True


def main():
    if not POSTS_DIR.is_dir():
        raise SystemExit("找不到 _posts 目录：{}".format(POSTS_DIR))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    rows = collect()

    csv_path = OUT_DIR / CSV_NAME
    write_csv(rows, csv_path)

    print("已导出 {} 篇文章 →".format(len(rows)))
    print("  {}".format(csv_path))

    xlsx_path = OUT_DIR / "文章分类标签表.xlsx"
    if write_xlsx(rows, xlsx_path):
        print("  {}（需要 openpyxl，已装）".format(xlsx_path))
    elif xlsx_path.exists():
        print("  （检测到旧的 .xlsx 可能已过期，本次未更新——未安装 openpyxl）")

    # 顺便给个分布小结，方便一眼看出有没有漏填
    from collections import Counter
    c1 = Counter(r["一级分类"] or "(空)" for r in rows)
    c2 = Counter((r["一级分类"] or "(空)", r["二级分类"] or "(空)") for r in rows)
    print("\n一级分类分布：" + "，".join("{} {} 篇".format(k, v) for k, v in c1.most_common()))
    print("二级分类共 {} 个：".format(len(c2)))
    for (a, b), n in sorted(c2.items(), key=lambda kv: (-kv[1], kv[0])):
        print("  {}/{}  {} 篇".format(a, b, n))

    empty = [r["文件名"] for r in rows if not r["一级分类"]]
    if empty:
        print("\n⚠ 有 {} 篇没有一级分类：".format(len(empty)))
        for f in empty:
            print("   - " + f)


if __name__ == "__main__":
    main()
